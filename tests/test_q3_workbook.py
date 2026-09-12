"""Workbook-payload tests independent of the XLSX authoring engine."""
import sys
import unittest
from pathlib import Path

sys.path[:0] = [str(Path(__file__).resolve().parents[1] / p)
                for p in ('src',)]
import pandas as pd
from openpyxl import load_workbook

from q3_workbook_data import build_payload, contiguous_emergency_ranges


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / 'C题_标准化数据与修正模板/C_final_delivery/templates/result3.xlsx'


class Q3WorkbookTests(unittest.TestCase):
    def test_contiguous_emergency_slots_are_merged_and_summed(self):
        """Separating adjacent slots would overstate the number of emergency events."""
        frame = pd.DataFrame({'interval_id': [1, 2, 4], 'e': [2., 3., 7.]})
        self.assertEqual(contiguous_emergency_ranges(frame), [
            {'start': 1, 'end': 3, 'kwh': 5.},
            {'start': 4, 'end': 5, 'kwh': 7.},
        ])

    def test_zero_and_tiny_numerical_emergency_values_are_omitted(self):
        """Solver noise below the physical reporting tolerance is not an event."""
        frame = pd.DataFrame({'interval_id': [0, 1, 2],
                              'e': [0., 1e-9, 2.]})
        self.assertEqual(contiguous_emergency_ranges(frame), [
            {'start': 2, 'end': 3, 'kwh': 2.},
        ])

    def test_template_shape_is_fixed(self):
        """Writing against a changed template must fail before cell offsets drift."""
        workbook = load_workbook(TEMPLATE, read_only=False)
        self.assertEqual([(workbook[s].max_row, workbook[s].max_column)
                          for s in workbook.sheetnames],
                         [(335, 147), (335, 147), (26, 6), (11, 3)])

    def test_formal_payload_contains_every_day_and_every_emergency_event(self):
        """The ellipsis rows in the template must expand to the complete 334-day result."""
        dispatch_path = ROOT / 'outputs/q3_final/dispatch_10min.csv'
        if not dispatch_path.exists():
            self.skipTest('formal causal backtest has not run yet')
        payload = build_payload(dispatch_path)
        self.assertEqual(len(payload['plan'][0]), 146)
        self.assertEqual(len(payload['adjustment'][0]), 146)
        self.assertEqual(len(payload['battery']), 334 * 6)
        self.assertEqual(payload['battery'][0][0], '2025-02-01')
        self.assertEqual(payload['battery'][-6][0], '2025-12-31')
        dated_rows = [row for row in payload['emergency'] if row[0] is not None]
        self.assertEqual(len(dated_rows), 334)
        dispatch = pd.read_csv(dispatch_path)
        first = dispatch[dispatch.date == '2025-02-01']
        self.assertAlmostEqual(payload['plan'][0][-2], float(first.g.sum()))
        self.assertAlmostEqual(payload['plan'][0][-1], float((first.price*first.g).sum()))
        self.assertAlmostEqual(sum(row[2] for row in payload['emergency']),
                               float(dispatch.e.sum()), places=6)


if __name__ == '__main__':
    unittest.main()
