"""Independent exported-workbook, physical and midnight reconstruction checks."""
import json,hashlib,re
from pathlib import Path
from q2_review_experiments import *
from q2_analyze import intervals
import openpyxl

def check_paper_dates(text,specified):
    for date,s in specified.items():
        section=text.split('### '+date,1)[1].split('### ',1)[0]
        for key in ('plan_kwh','plan_cost','emergency_kwh','emergency_cost','total_cost','initial','terminal'):
            if key in s: assert f'{s[key]:,.2f}' in section,(date,key)
        if 'specified_g' in s:
            for v in s['specified_g']: assert f'{v:,.2f}' in section
            for b in s['blocks']:
                assert b[0] in section
                assert f'{b[1]:,.2f}' in section and f'{b[2]:,.2f}' in section
    dates=[d for d,s in specified.items() if s['emergency_kwh']>EPS]
    phrase='、'.join(dates) if dates else '无'
    assert f'发生紧急购电的日期为：{phrase}；' in text

def main():
    protocol=json.loads((OUT/'selection_protocol.json').read_text())
    final=json.loads((OUT/'experiments'/f"{protocol['final_policy']}.json").read_text())
    config=Config(**final['config'])
    x=pd.read_csv(OUT/'dispatch_10min.csv'); daily=pd.read_csv(OUT/'daily_summary.csv')
    physical=audit(daily,x,config); payload=json.loads((OUT/'workbook_payload.json').read_text(encoding='utf8'))
    wb=openpyxl.load_workbook(OUT/'result2.xlsx',data_only=True)
    assert wb.sheetnames==['计划购电量','充放电量','紧急购电量']
    plan=wb.worksheets[0]; cells=0
    for j,(date,z) in enumerate(x.groupby('date',sort=True),2):
        assert str(plan.cell(j,1).value)[:10]==date
        expected=np.r_[z.g,z.g.sum(),z.price@z.g]
        actual=np.array([plan.cell(j,k).value for k in range(2,148)],float)
        assert np.max(abs(expected-actual))<1e-7,(date,'plan')
        cells+=len(expected)
    assert plan.max_row==335 and plan.max_column==147
    for ws,rows in zip(wb.worksheets[1:],[payload['battery'],payload['emergency']]):
        for r,values in enumerate(rows,2):
            for c,v in enumerate(values,1):
                got=ws.cell(r,c).value
                if v is None: assert got is None,(ws.title,r,c)
                elif c==1: assert str(got)[:10]==v
                elif isinstance(v,(int,float)): assert abs(got-v)<1e-7
                else: assert got==v
        assert ws.max_row==len(rows)+1
    # Derive summaries independently from the detail, not only payload matching.
    for day,(date,z) in enumerate(x.groupby('date',sort=True)):
        for k in range(6):
            row=2+day*6+k; a=z.iloc[24*k:24*(k+1)]
            assert abs(wb.worksheets[1].cell(row,3).value-a.c.sum())<1e-6
            assert abs(wb.worksheets[1].cell(row,4).value-a.d.sum())<1e-6
        assert abs(wb.worksheets[1].cell(2+day*6,6).value-z.E_start.iloc[0])<1e-6
        assert abs(wb.worksheets[1].cell(3+day*6,6).value-z.E_end.iloc[-1])<1e-6
    er=[]
    for date,z in x.groupby('date',sort=True):
        events=intervals(z)
        er.extend([[date if j==0 else None,*v] for j,v in enumerate(events)] if events else [[date,'无',0.]])
    assert len(er)==len(payload['emergency'])
    for expected,got in zip(er,payload['emergency']):
        assert expected[:2]==got[:2] and abs(expected[2]-got[2])<1e-7
    assert all(cell.data_type!='e' for s in wb for row in s for cell in row)
    # Replay final policy from available midnight information and inherited state.
    inputs=load_inputs(); l,v,p,forecasts=inputs
    rotation=pd.read_csv(OUT/'rolling_selection_T0.csv')
    assert (rotation.last_history_day<rotation.decision_day).all()
    assert (rotation.end_exclusive==rotation.decision_day).all()
    for row in rotation.itertuples():
        scores={m:getattr(row,m) for m in ('M1','M3','M4')}
        assert row.winner==min(scores,key=lambda m:(scores[m],m))
    errors=[]; actual_errors=[]; state=6000.
    for row in daily.itertuples():
        month=int(row.date[5:7]); model='M1' if month==2 else rotation.loc[rotation.month==month,'winner'].iloc[0]
        assert model==row.model
        f=forecasts['P0' if model=='M1' else 'P1']; pn=f[:,0]-f[:,1]
        net=planning_input(l-v,pn,row.day,model,.8,14)
        opt=optimize(net,p,state,target=6000.)
        z=x[x.date==row.date]
        errors.append(float(np.max(abs(opt['g']-z.g.to_numpy()))))
        r=execute_closed(opt['g'],l[row.day],v[row.day],state,6000.) if row.day==364 else execute(opt['g'],l[row.day],v[row.day],state)
        for key in ('c','d','e','unused','spill'): actual_errors.append(float(np.max(abs(r[key]-z[key].to_numpy()))))
        state=float(r['E'][-1])
    assert max(errors)<1e-6 and max(actual_errors)<1e-6
    # Fresh all-experiment audit, including nonstandard initial inventories.
    count=0
    for path in (OUT/'experiments').glob('*.json'):
        record=json.loads(path.read_text()); stem=path.stem
        d=pd.read_csv(path.parent/f'{stem}_daily.csv'); z=pd.read_csv(path.parent/f'{stem}_dispatch.csv')
        audit(d,z,Config(**record['config'])); count+=1
    paper=ROOT/'docs/result/第二问求解与结果分析.md'; text=paper.read_text(encoding='utf8')
    check_paper_dates(text,payload['specified'])
    assert text.count('$$')%2==0
    # Links are checked from the manuscript directory so the package remains
    # valid after cloning or extracting it on another machine.
    for target in re.findall(r'\]\(([^)]+)\)', text):
        target = target.split('#', 1)[0]
        if not target or target.startswith(('http://', 'https://', 'mailto:')):
            continue
        path = (paper.parent / target).resolve()
        assert not Path(target).is_absolute(), target
        assert path.exists(), target
    expected_cost=float((x.price*x.g+5*x.price*x.e).sum())
    assert f'{expected_cost:,.2f}' in text
    code=list((ROOT/'src').glob('q2*.py'))+[ROOT/'src/q2_workbook.mjs']
    sources=[DATA/'data_csv/annual_actual_10min.csv',DATA/'data_csv/typical_day_10min.csv',OLD/'forecasts.npz']+code
    hashes={str(path.relative_to(ROOT)):hashlib.sha256(path.read_bytes()).hexdigest() for path in sources}
    (OUT/'source_hashes.json').write_text(json.dumps(hashes,indent=2),encoding='utf8')
    result=dict(selected=protocol['final_policy'],days=334,dispatch_slots=len(x),plan_numeric_cells=cells,
        experiments_audited=count,physical=physical,maximum_midnight_plan_error_kwh=max(errors),
        maximum_execution_reproduction_error_kwh=max(actual_errors),cost=expected_cost,
        emergency_kwh=float(x.e.sum()),emergency_events=sum(v[2]>EPS for v in er),
        paper_sha256=hashlib.sha256(paper.read_bytes()).hexdigest(),xlsx_sha256=hashlib.sha256((OUT/'result2.xlsx').read_bytes()).hexdigest())
    (OUT/'delivery_verification.json').write_text(json.dumps(result,indent=2),encoding='utf8')
    print(json.dumps(result,indent=2))

if __name__=='__main__':main()
