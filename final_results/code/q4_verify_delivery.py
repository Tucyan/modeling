"""Independent physical, accounting and workbook verification for Q4."""
from __future__ import annotations

from pathlib import Path
import argparse
import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from q4_solver import settle


ROOT = Path(__file__).resolve().parents[1]
DATES = pd.date_range('2025-02-01', '2025-12-31').strftime('%Y-%m-%d').tolist()
EPS = 1e-6


def _finite(frame, name):
    values = frame.select_dtypes(include=[np.number]).to_numpy(dtype=float)
    if not np.isfinite(values).all():
        raise AssertionError(f'{name} contains non-finite numeric values')


def _verify_ledger(path, mode, settlement):
    dispatch = pd.read_csv(path)
    required = {'date', 'interval_id', 'g', 'q', 'c', 'd', 'e', 'unused', 'spill',
                'E_start', 'E_end', 'load', 'pv', 'price', 'plan_price', 'delta'}
    if required - set(dispatch.columns):
        raise AssertionError(f'{path} missing {sorted(required-set(dispatch.columns))}')
    dispatch['date'] = dispatch.date.astype(str)
    if len(dispatch) != 334 * 144 or dispatch.duplicated(['date', 'interval_id']).any():
        raise AssertionError(f'{path} is not a complete 334-day ledger')
    if dispatch.date.unique().tolist() != DATES:
        raise AssertionError(f'{path} dates are incomplete or unordered')
    if not np.array_equal(dispatch.interval_id.to_numpy(), np.tile(np.arange(144), 334)):
        raise AssertionError(f'{path} interval order is invalid')
    _finite(dispatch, str(path))
    if (dispatch[['g', 'q', 'c', 'd', 'e', 'unused', 'spill']] < -EPS).any().any():
        raise AssertionError(f'{path} contains negative flows')
    if (dispatch.E_start < 1200-EPS).any() or (dispatch.E_end < 1200-EPS).any():
        raise AssertionError(f'{path} violates lower SOC bound')
    if (dispatch.E_start > 10800+EPS).any() or (dispatch.E_end > 10800+EPS).any():
        raise AssertionError(f'{path} violates upper SOC bound')
    if dispatch.c.max() > 5000/6 + EPS or dispatch.d.max() > 5000/6 + EPS:
        raise AssertionError(f'{path} violates power bound')
    bus = dispatch.q - dispatch.unused + dispatch.e + dispatch.pv - dispatch.spill + dispatch.d - dispatch.load - dispatch.c
    state = dispatch.E_end - dispatch.E_start - .9*dispatch.c + dispatch.d/.9
    if float(np.max(np.abs(bus))) > EPS or float(np.max(np.abs(state))) > EPS:
        raise AssertionError(f'{path} physical residual exceeds tolerance')
    if (dispatch.unused > dispatch.q + EPS).any() or (dispatch.spill > dispatch.pv + EPS).any():
        raise AssertionError(f'{path} disposal bound violated')
    if np.any((dispatch.c > EPS) & (dispatch.d > EPS)):
        raise AssertionError(f'{path} simultaneous charge/discharge')
    if mode == 'q4-2' and not np.allclose(dispatch.g, dispatch.q, atol=EPS, rtol=0):
        raise AssertionError('Q4-2 must keep q equal to g')
    recomputed = settle(dispatch.g.to_numpy(), dispatch.q.to_numpy(),
                        dispatch.price.to_numpy(), settlement) + 5 * dispatch.price * dispatch.e
    total = float(recomputed.sum())
    daily = dispatch.assign(_cost=recomputed).groupby('date', sort=False)._cost.sum()
    if abs(total - float(daily.sum())) > EPS:
        raise AssertionError('daily and ledger costs disagree')
    if abs(float(dispatch.E_end.iloc[-1]) - 6000) > EPS:
        raise AssertionError('final SOC is not 6000')
    if not np.allclose(dispatch.E_end.to_numpy()[:-1], dispatch.E_start.to_numpy()[1:], atol=EPS, rtol=0):
        raise AssertionError('cross-day SOC discontinuity')
    return dict(rows=len(dispatch), days=334, cost=total,
                emergency_kwh=float(dispatch.e.sum()),
                max_bus_residual=float(np.max(np.abs(bus))),
                max_state_residual=float(np.max(np.abs(state))),
                final_soc=float(dispatch.E_end.iloc[-1]))


def _verify_versions(path, dispatch, mode):
    versions = pd.read_csv(path)
    if mode == 'q4-2':
        if len(versions) != 0:
            raise AssertionError('Q4-2 contains contract updates')
        return 0
    if len(versions) and not set(versions.issue_hour.astype(int)).issubset({6, 12, 18}):
        raise AssertionError('invalid Q4-3 update hour')
    if len(versions):
        if (versions.interval_id < versions.issue_hour * 6).any():
            raise AssertionError('version rewrites an already executed slot')
        for date, group in versions.groupby('date', sort=False):
            day = dispatch[dispatch.date == str(date)].sort_values('interval_id')
            reconstructed = day.g.to_numpy(dtype=float)
            for row in group.sort_values(['issue_hour', 'interval_id']).itertuples():
                reconstructed[int(row.interval_id)] = float(row.q)
            np.testing.assert_allclose(reconstructed, day.q.to_numpy(dtype=float), atol=EPS, rtol=0)
    return int(len(versions))


def _verify_forecasts(path, mode):
    frame = pd.read_csv(path)
    required = {'date', 'day', 'issue_hour', 'interval_id', 'forecast_price',
                'actual_price', 'history_start_day', 'history_end_exclusive',
                'observed_prefix_slots'}
    if required - set(frame.columns):
        raise AssertionError(f'{path} missing forecast columns')
    expected_rows = 334 * (144 if mode == 'q4-2' else 360)
    if len(frame) != expected_rows or frame.duplicated(['date', 'issue_hour', 'interval_id']).any():
        raise AssertionError(f'{path} has incomplete forecast records')
    _finite(frame, str(path))
    if not set(frame.issue_hour.astype(int)).issubset({0, 6, 12, 18}):
        raise AssertionError(f'{path} has invalid issue hour')
    if (frame.history_end_exclusive > frame.day).any():
        raise AssertionError(f'{path} uses future history')
    if (frame.history_start_day < 7).any():
        raise AssertionError(f'{path} uses unavailable warm-up history')
    if (frame.observed_prefix_slots != frame.issue_hour * 6).any():
        raise AssertionError(f'{path} has invalid observed prefix')
    if (frame.forecast_price <= 0).any() or (frame.actual_price <= 0).any():
        raise AssertionError(f'{path} contains non-positive price')
    return int(len(frame))


def _verify_workbook(path, mode):
    template_path = ROOT / 'C题_标准化数据与修正模板/C_final_delivery/templates' / (
        'result4-2.xlsx' if mode == 'q4-2' else 'result4-3.xlsx')
    expected = load_workbook(template_path, data_only=False, read_only=True).sheetnames
    workbook = load_workbook(path, data_only=False, read_only=False)
    if workbook.sheetnames != expected:
        raise AssertionError(f'{path} sheets differ from template')
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, float) and not np.isfinite(cell.value):
                    raise AssertionError(f'{path} has non-finite cell')
    if workbook.worksheets[0].max_row < 335 or workbook.worksheets[0].max_column < 147:
        raise AssertionError('plan sheet is truncated')
    if mode == 'q4-3' and (workbook.worksheets[1].max_row < 335
                            or workbook.worksheets[1].max_column < 147):
        raise AssertionError('adjustment sheet is truncated')
    return dict(sheets=workbook.sheetnames, plan_shape=[335, 147])


def verify(output):
    output = Path(output)
    result = {}
    for settlement in ('N0', 'N1'):
        folder = output / settlement
        for mode in ('q4-2', 'q4-3'):
            dispatch_path = folder / f'{mode}_dispatch_10min.csv'
            versions_path = folder / f'{mode}_contract_versions.csv'
            ledger = _verify_ledger(dispatch_path, mode, settlement)
            dispatch = pd.read_csv(dispatch_path)
            dispatch['date'] = dispatch.date.astype(str)
            ledger['versions'] = _verify_versions(versions_path, dispatch, mode)
            ledger['forecast_records'] = _verify_forecasts(
                folder / f'{mode}_forecast_records.csv', mode)
            workbook_path = folder / f'result4-{"2" if mode == "q4-2" else "3"}.xlsx'
            ledger['workbook'] = _verify_workbook(workbook_path, mode)
            result[f'{settlement}/{mode}'] = ledger
    (output / 'delivery_verification.json').write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', default=str(ROOT / 'outputs/q4_final'))
    verify(parser.parse_args().output)
