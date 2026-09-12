"""Independent disk-level verification of the formal Question 3 delivery."""
from __future__ import annotations

from pathlib import Path
import hashlib
import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs/q3_final'
EPS = 1e-6
CAP = 5000 / 6


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _assert_small(checks, tolerance=EPS):
    bad = {name: value for name, value in checks.items()
           if not np.isfinite(float(value)) or float(value) > tolerance}
    if bad:
        raise AssertionError(f'Verification residuals exceed tolerance: {bad}')


def _reconstruct_contracts(dispatch, versions):
    reconstructed = np.empty(len(dispatch), dtype=float)
    for _, day in dispatch.groupby('date', sort=True):
        day = day.sort_values('interval_id')
        current = day.g.to_numpy(dtype=float).copy()
        records = versions[versions.date == day.iloc[0].date] if len(versions) else versions
        for issue_hour, issue in records.groupby('issue_hour', sort=True):
            issue = issue.sort_values('interval_id')
            start = int(issue_hour) * 6
            if list(issue.interval_id.astype(int)) != list(range(start, 144)):
                raise AssertionError('A contract version does not cover its complete remaining horizon')
            previous = current[start:].copy()
            original = day.g.to_numpy(dtype=float)[start:]
            _assert_small({
                'version_original': float(np.max(abs(issue.original_g.to_numpy()-original))),
                'version_previous': float(np.max(abs(issue.previous_q.to_numpy()-previous))),
                'version_delta': float(np.max(abs(issue.delta.to_numpy()-(issue.q.to_numpy()-original)))),
                'version_change': float(np.max(abs(issue.change.to_numpy()-(issue.q.to_numpy()-previous)))),
            })
            current[start:] = issue.q.to_numpy(dtype=float)
        reconstructed[day.index.to_numpy()] = current
    return reconstructed


def verify_dispatch(daily, dispatch, versions, settlement='N0',
                    expected_start='2025-02-01', expected_end='2025-12-31'):
    if settlement not in ('N0', 'N1'):
        raise ValueError('Supported settlements are N0 and N1')
    required = {'date', 'interval_id', 'g', 'q', 'c', 'd', 'e', 'unused',
                'spill', 'E_start', 'E_end', 'load', 'pv', 'price'}
    if required - set(dispatch.columns):
        raise AssertionError(f'Missing dispatch fields: {sorted(required-set(dispatch.columns))}')
    expected_dates = pd.date_range(expected_start, expected_end).strftime('%Y-%m-%d').tolist()
    if len(daily) != len(expected_dates) or len(dispatch) != len(expected_dates) * 144:
        raise AssertionError('Delivery has incomplete dates or slots')
    if 'date' not in daily.columns or daily.date.astype(str).tolist() != expected_dates:
        raise AssertionError('Daily dates must be complete and ordered')
    if dispatch.date.astype(str).nunique() != len(expected_dates):
        raise AssertionError('Dispatch dates are incomplete')
    if dispatch.date.astype(str).tolist() != list(np.repeat(expected_dates, 144)):
        raise AssertionError('Dispatch dates must be complete and ordered')
    numeric_columns = [column for column in required if column not in ('date', 'interval_id')]
    if not np.isfinite(dispatch[numeric_columns].to_numpy(dtype=float)).all():
        raise AssertionError('Dispatch contains non-finite numeric values')
    version_required = {'date','issue_hour','interval_id','original_g',
                        'previous_q','q','delta','change'}
    if version_required - set(versions.columns):
        raise AssertionError('Contract versions are missing required fields')
    if len(versions):
        numeric_versions = versions[list(version_required-{'date'})].to_numpy(dtype=float)
        if not np.isfinite(numeric_versions).all():
            raise AssertionError('Contract versions contain non-finite numeric values')
        if versions.date.isna().any() or not set(versions.date.astype(str)).issubset(expected_dates):
            raise AssertionError('Contract versions contain unknown dates')
        if versions.duplicated(['date','issue_hour','interval_id']).any():
            raise AssertionError('Duplicate contract version keys')
        issue_hours = versions.issue_hour.to_numpy(dtype=float)
        if (not np.equal(issue_hours, np.floor(issue_hours)).all()
                or not set(issue_hours.astype(int)).issubset({0, 6, 12, 18})):
            raise AssertionError('Contract versions contain an unsupported issue hour')
    if 'cost' not in daily.columns or not np.isfinite(daily.cost.to_numpy(dtype=float)).all():
        raise AssertionError('Daily summary contains non-finite cost values')
    if dispatch.duplicated(['date', 'interval_id']).any():
        raise AssertionError('Duplicate dispatch keys')
    if not np.array_equal(dispatch.interval_id.to_numpy(), np.tile(np.arange(144), len(expected_dates))):
        raise AssertionError('Incorrect interval order')
    reconstructed = _reconstruct_contracts(dispatch, versions)
    down_coefficient = .5 if settlement == 'N0' else -.5
    regular = dispatch.price * (
        dispatch.g + 1.5*np.maximum(dispatch.q-dispatch.g, 0)
        + down_coefficient*np.maximum(dispatch.g-dispatch.q, 0))
    recomputed_cost = float(regular.sum() + (5*dispatch.price*dispatch.e).sum())
    residuals = {
        'bus': float(abs(dispatch.q-dispatch.unused+dispatch.e+dispatch.pv
                         -dispatch.spill+dispatch.d-dispatch.load-dispatch.c).max()),
        'state': float(abs(dispatch.E_end-dispatch.E_start-.9*dispatch.c
                           +dispatch.d/.9).max()),
        'continuity': float(abs(dispatch.E_end.to_numpy()[:-1]
                                -dispatch.E_start.to_numpy()[1:]).max()),
        'bounds': float(max(0, 1200-dispatch.E_start.min(), 1200-dispatch.E_end.min(),
                            dispatch.E_start.max()-10800, dispatch.E_end.max()-10800)),
        'power': float(max(0, dispatch.c.max()-CAP, dispatch.d.max()-CAP)),
        'disposal': float(max(0, (dispatch.unused-dispatch.q).max(),
                               (dispatch.spill-dispatch.pv).max())),
        'negative': float(max(0, -dispatch[['g','q','c','d','e','unused','spill']].min().min())),
        'contract_reconstruction': float(np.max(abs(reconstructed-dispatch.q.to_numpy()))),
    }
    overlap = int(((dispatch.c > EPS) & (dispatch.d > EPS)).sum())
    cost_error = abs(recomputed_cost - float(daily.cost.sum()))
    _assert_small(residuals)
    if overlap:
        raise AssertionError(f'{overlap} simultaneous charge/discharge slots')
    _assert_small({'cost': cost_error})
    if abs(float(dispatch.E_start.iloc[0])-6000) > EPS or abs(float(dispatch.E_end.iloc[-1])-6000) > EPS:
        raise AssertionError('Evaluation boundary SOC is not 6000 kWh')
    return {'rows': len(dispatch), 'days': len(daily),
            'max_residuals_kwh': residuals, 'overlap_slots': overlap,
            'cost': recomputed_cost, 'cost_error_yuan': cost_error,
            'all_contract_versions_reconstructed': True}


def _date_text(value):
    return pd.Timestamp(value).strftime('%Y-%m-%d') if value is not None else None


def _expected_emergency(dispatch):
    rows = []
    for date, day in dispatch.groupby('date', sort=True):
        positive = day[day.e > 1e-7].sort_values('interval_id')
        events = []
        for record in positive.itertuples(index=False):
            slot = int(record.interval_id)
            if events and slot == events[-1]['end']:
                events[-1]['end'] += 1
                events[-1]['kwh'] += float(record.e)
            else:
                events.append({'start': slot, 'end': slot+1, 'kwh': float(record.e)})
        if not events:
            rows.append([date, '无', 0.])
        for index, event in enumerate(events):
            def stamp(boundary):
                minutes = boundary * 10
                return f'{minutes//60:02d}:{minutes%60:02d}'
            rows.append([date if index == 0 else None,
                         f"{stamp(event['start'])}-{stamp(event['end'])}", event['kwh']])
    return rows


def verify_workbook(workbook_path, dispatch, settlement='N0'):
    if settlement not in ('N0', 'N1'):
        raise ValueError('Supported settlements are N0 and N1')
    workbook = load_workbook(workbook_path, data_only=True, read_only=False)
    if workbook.sheetnames != ['计划购电量', '调整购电量', '充放电量', '紧急购电量']:
        raise AssertionError('Workbook sheet structure changed')
    g = dispatch.g.to_numpy().reshape(334, 144)
    q = dispatch.q.to_numpy().reshape(334, 144)
    plan_sheet, adjust_sheet = workbook['计划购电量'], workbook['调整购电量']
    plan = np.asarray([[plan_sheet.cell(row, col).value for col in range(2, 146)]
                       for row in range(2, 336)], dtype=float)
    adjustment = np.asarray([[adjust_sheet.cell(row, col).value for col in range(2, 146)]
                             for row in range(2, 336)], dtype=float)
    if not np.isfinite(plan).all() or not np.isfinite(adjustment).all():
        raise AssertionError('Workbook plan/adjustment matrix contains non-finite values')
    checks = {'plan_matrix': float(np.max(abs(plan-g))),
              'adjustment_matrix': float(np.max(abs(adjustment-q)))}
    for index, (_, day) in enumerate(dispatch.groupby('date', sort=True), start=2):
        plan_total = float(plan_sheet.cell(index, 146).value)
        plan_cost = float(plan_sheet.cell(index, 147).value)
        adjustment_total = float(adjust_sheet.cell(index, 146).value)
        adjustment_cost = float(adjust_sheet.cell(index, 147).value)
        if not np.isfinite((plan_total, plan_cost, adjustment_total, adjustment_cost)).all():
            raise AssertionError('Workbook totals contain non-finite values')
        down = .5 if settlement == 'N0' else -.5
        regular = day.price * (day.g + 1.5*np.maximum(day.q-day.g, 0)
                               + down*np.maximum(day.g-day.q, 0))
        checks['plan_total'] = max(checks.get('plan_total', 0), abs(plan_total-day.g.sum()))
        checks['plan_cost'] = max(checks.get('plan_cost', 0), abs(plan_cost-(day.price*day.g).sum()))
        checks['adjustment_total'] = max(checks.get('adjustment_total', 0), abs(adjustment_total-day.q.sum()))
        checks['adjustment_cost'] = max(checks.get('adjustment_cost', 0), abs(adjustment_cost-regular.sum()))
    battery = workbook['充放电量']
    row = 2
    for date, day in dispatch.groupby('date', sort=True):
        for block in range(6):
            part = day.iloc[block*24:(block+1)*24]
            expected = (float(part.c.sum()), float(part.d.sum()))
            charge = float(battery.cell(row, 3).value)
            discharge = float(battery.cell(row, 4).value)
            if not np.isfinite((charge, discharge)).all():
                raise AssertionError('Workbook battery table contains non-finite values')
            checks['battery_charge'] = max(checks.get('battery_charge', 0),
                                            abs(charge-expected[0]))
            checks['battery_discharge'] = max(checks.get('battery_discharge', 0),
                                               abs(discharge-expected[1]))
            if block == 0 and _date_text(battery.cell(row, 1).value) != date:
                raise AssertionError('Battery date mismatch')
            row += 1
    if row != 2006:
        raise AssertionError('Battery row count mismatch')
    if any(battery.cell(tail, col).value is not None
           for tail in range(row, battery.max_row+1) for col in range(1, 5)):
        raise AssertionError('Battery table contains trailing rows')
    emergency_sheet = workbook['紧急购电量']
    expected_rows = _expected_emergency(dispatch)
    actual_rows = []
    for index in range(2, 2+len(expected_rows)):
        kwh = float(emergency_sheet.cell(index, 3).value)
        if not np.isfinite(kwh):
            raise AssertionError('Workbook emergency table contains non-finite values')
        actual_rows.append([_date_text(emergency_sheet.cell(index, 1).value),
                            emergency_sheet.cell(index, 2).value, kwh])
    for expected, actual in zip(expected_rows, actual_rows):
        if expected[:2] != actual[:2]:
            raise AssertionError(f'Emergency event mismatch: {expected[:2]} != {actual[:2]}')
        checks['emergency_kwh'] = max(checks.get('emergency_kwh', 0),
                                      abs(expected[2]-actual[2]))
    emergency_end = 2 + len(expected_rows)
    if any(emergency_sheet.cell(tail, col).value is not None
           for tail in range(emergency_end, emergency_sheet.max_row+1)
           for col in range(1, 4)):
        raise AssertionError('Emergency table contains trailing rows')
    errors = []
    for sheet in workbook:
        for cells in sheet.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str) and cell.value.startswith(('#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A', '#NUM!', '#NULL!', '#SPILL!', '#CALC!')):
                    errors.append(f'{sheet.title}!{cell.coordinate}:{cell.value}')
    if errors:
        raise AssertionError(f'Workbook formula errors: {errors[:10]}')
    _assert_small(checks)
    return {'sheet_dimensions': {sheet.title: sheet.calculate_dimension() for sheet in workbook},
            'maximum_cell_error': max(checks.values()), 'checks': checks,
            'formula_errors': 0, 'emergency_rows': len(expected_rows)}


def verify_delivery(root=ROOT):
    root = Path(root)
    out = root / 'outputs/q3_final'
    daily = pd.read_csv(out / 'daily_summary.csv')
    dispatch = pd.read_csv(out / 'dispatch_10min.csv')
    versions = pd.read_csv(out / 'contract_versions.csv')
    protocol = json.loads((out / 'selection_protocol.json').read_text(encoding='utf-8'))
    if protocol['last_history_day'] >= protocol['decision_day']:
        raise AssertionError('Protocol contains future selection data')
    dispatch_checks = verify_dispatch(daily, dispatch, versions, protocol['settlement'])
    workbook_checks = verify_workbook(out / 'result3.xlsx', dispatch,
                                      protocol['settlement'])
    q2_record = json.loads((root / 'outputs/q2_review_v2/delivery_verification.json').read_text())
    q2_paper = root / 'docs/result/第二问求解与结果分析.md'
    q2_workbook = root / 'outputs/q2_review_v2/result2.xlsx'
    q2_checks = {
        'paper_expected': q2_record['paper_sha256'], 'paper_actual': sha256(q2_paper),
        'xlsx_expected': q2_record['xlsx_sha256'], 'xlsx_actual': sha256(q2_workbook),
    }
    if q2_checks['paper_expected'] != q2_checks['paper_actual'] or q2_checks['xlsx_expected'] != q2_checks['xlsx_actual']:
        raise AssertionError('Q2 formal delivery changed')
    hashed = [
        root / 'C题_标准化数据与修正模板/C_final_delivery/data_csv/annual_actual_10min.csv',
        root / 'C题_标准化数据与修正模板/C_final_delivery/data_csv/pv_forecast_hourly.csv',
        root / 'C题_标准化数据与修正模板/C_final_delivery/data_csv/typical_day_10min.csv',
        root / 'src/q3_solver.py', root / 'src/q3_forecast.py',
        root / 'src/q3_experiments.py', root / 'src/q3_workbook_data.py',
        root / 'src/q3_workbook.mjs', root / 'src/q3_report.py',
        out / 'daily_summary.csv', out / 'dispatch_10min.csv',
        out / 'contract_versions.csv', out / 'comparison.csv',
        out / 'selection_protocol.json', out / 'result3.xlsx',
        root / 'docs/result/第三问求解与结果分析.md',
    ]
    result = dict(
        selected=protocol['selected'], selection_winner=protocol['winner'],
        selection_is_strictly_historical=True, dispatch=dispatch_checks,
        workbook=workbook_checks, q2_unchanged=True, q2_hashes=q2_checks,
        sha256={str(path.relative_to(root)): sha256(path) for path in hashed})
    (out / 'delivery_verification.json').write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    return result


def verify_unified_delivery(output, root=ROOT):
    """Verify both unified settlement deliveries and all saved ablations."""
    output = Path(output).resolve(); root = Path(root).resolve()
    stage_path = output / 'stage_unified-ablations.json'
    if not stage_path.is_file():
        raise AssertionError('Unified ablation stage manifest is missing')
    stage = json.loads(stage_path.read_text(encoding='utf-8'))
    if stage.get('status') != 'complete':
        raise AssertionError('Unified ablation stage is incomplete')
    for relative, expected in stage.get('artifacts', {}).items():
        artifact = output / relative
        if not artifact.is_file() or sha256(artifact) != expected:
            raise AssertionError(f'Unified stage artifact is stale: {relative}')
    results = {'output': str(output), 'stage_manifest_verified': True,
               'settlements': {}}
    for settlement in ('N0', 'N1'):
        folder = output / settlement
        protocol = json.loads((folder / 'selection_protocol.json').read_text(encoding='utf-8'))
        if protocol['settlement'] != settlement:
            raise AssertionError('Protocol settlement mismatch')
        if protocol['last_history_day'] >= protocol['decision_day'] or len(protocol['scores']) != 48:
            raise AssertionError('Protocol is not a complete strictly historical selection')
        daily = pd.read_csv(folder / 'daily_summary.csv')
        dispatch = pd.read_csv(folder / 'dispatch_10min.csv')
        versions = pd.read_csv(folder / 'contract_versions.csv')
        main = verify_dispatch(daily, dispatch, versions, settlement)
        workbook = verify_workbook(folder / 'result3.xlsx', dispatch, settlement)
        policies = {}
        selected_label = f"calibrated_{protocol['selected']['release_name']}"
        for label in (f'calibrated_{name}' for name in
                      ('R0','R6','R12','R18','R6_12','R6_18','R12_18','Rall')):
            base = folder if label == selected_label else folder / 'release_ablations'
            d = pd.read_csv(base / f'{label}_daily.csv')
            x = pd.read_csv(base / f'{label}_dispatch.csv')
            v = pd.read_csv(base / f'{label}_versions.csv')
            policies[label] = verify_dispatch(d, x, v, settlement)
        comparison = pd.read_csv(folder / 'comparison.csv').set_index('label')
        expected_comparison = set(policies) | {'Q2_frozen'}
        if set(comparison.index) != expected_comparison:
            raise AssertionError('Comparison table does not match release ablations')
        _assert_small({f'comparison_{label}':
                       abs(float(comparison.loc[label, 'cost'])-check['cost'])
                       for label, check in policies.items()})
        label = 'Q2_day_ahead_with_updates'
        adapter_daily = pd.read_csv(folder / f'{label}_daily.csv')
        adapter_dispatch = pd.read_csv(folder / f'{label}_dispatch.csv')
        policies[label] = verify_dispatch(
            adapter_daily, adapter_dispatch,
            pd.read_csv(folder / f'{label}_versions.csv'), settlement)
        adapter_summary = json.loads((folder / f'{label}.json').read_text(encoding='utf-8'))
        _assert_small({'adapter_summary':
                       abs(adapter_summary['cost']-policies[label]['cost'])})
        fixed = json.loads((folder / 'fixed_ledger_repricing.json').read_text(encoding='utf-8'))
        for other in ('N0', 'N1'):
            down = .5 if other == 'N0' else -.5
            regular = dispatch.price * (
                dispatch.g + 1.5*np.maximum(dispatch.q-dispatch.g, 0)
                + down*np.maximum(dispatch.g-dispatch.q, 0))
            recomputed = float(regular.sum()+(5*dispatch.price*dispatch.e).sum())
            _assert_small({f'fixed_repricing_{other}': abs(fixed[other]-recomputed)})
        results['settlements'][settlement] = dict(
            selected=protocol['selected'], main=main, workbook=workbook,
            policies=policies)
    q2_daily = pd.read_csv(output / 'Q2_frozen_daily.csv')
    q2_dispatch = pd.read_csv(output / 'Q2_frozen_dispatch.csv')
    q2_versions = pd.read_csv(output / 'Q2_frozen_versions.csv')
    results['q2_frozen'] = {
        settlement: verify_dispatch(q2_daily, q2_dispatch, q2_versions, settlement)
        for settlement in ('N0', 'N1')}
    for settlement in ('N0', 'N1'):
        comparison = pd.read_csv(output/settlement/'comparison.csv').set_index('label')
        _assert_small({f'q2_comparison_{settlement}':
                       abs(float(comparison.loc['Q2_frozen', 'cost'])
                           - results['q2_frozen'][settlement]['cost'])})
    cross = json.loads((output/'cross_protocols.json').read_text(encoding='utf-8'))
    selected_same = (results['settlements']['N0']['selected']
                     == results['settlements']['N1']['selected'])
    if bool(cross.get('protocols_identical')) != selected_same:
        raise AssertionError('Cross-settlement protocol record is inconsistent')
    if selected_same and not cross.get('intentional_reuse'):
        raise AssertionError('Identical cross-settlement protocols were not explicitly recorded')
    if not selected_same and len(cross.get('evaluations', [])) != 2:
        raise AssertionError('Different protocols require two cross-settlement evaluations')
    results['cross_protocols'] = cross
    results['sha256'] = {
        str(path.relative_to(root)): sha256(path)
        for path in sorted(output.rglob('*')) if path.is_file()
        and path.suffix in ('.csv', '.json', '.xlsx')
        and path.name != 'delivery_verification.json'}
    (output / 'delivery_verification.json').write_text(
        json.dumps(results, ensure_ascii=False, indent=2), encoding='utf-8')
    return results


def main():
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--output')
    args = parser.parse_args()
    if args.output:
        result = verify_unified_delivery(args.output)
        print(json.dumps({settlement: {
            'selected': value['selected'], 'cost': value['main']['cost'],
            'maximum_cell_error': value['workbook']['maximum_cell_error']}
            for settlement, value in result['settlements'].items()}, ensure_ascii=False))
    else:
        result = verify_delivery()
        print(json.dumps({'selected': result['selected'],
                          'cost': result['dispatch']['cost'],
                          'maximum_cell_error': result['workbook']['maximum_cell_error'],
                          'q2_unchanged': result['q2_unchanged']}, ensure_ascii=False))


if __name__ == '__main__':
    main()
