"""Read-only independent Q4 artifact review; does not import delivery builders."""
from pathlib import Path
import hashlib
import json
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'outputs/q4_final'
actual = pd.read_csv(ROOT / 'C题_标准化数据与修正模板/C_final_delivery/data_csv/annual_actual_10min.csv').iloc[31*144:].reset_index(drop=True)
report = {}
def same(a, b):
    a, b = np.asarray(a, float), np.asarray(b, float)
    assert a.shape == b.shape
    assert np.isfinite(a).all() and np.isfinite(b).all()
    err = float(np.max(np.abs(a-b))) if a.size else 0.
    assert err < 1e-6, err
    return err

for settlement in ('N0', 'N1'):
    for mode in ('q4-2', 'q4-3'):
        folder = OUT / settlement
        df = pd.read_csv(folder / f'{mode}_dispatch_10min.csv')
        daily = pd.read_csv(folder / f'{mode}_daily_summary.csv')
        assert df.date.tolist() == actual.date.tolist()
        assert df.interval_id.tolist() == actual.interval_id.tolist()
        for col, source in [('load','load_kwh_10min'), ('pv','pv_actual_kwh_10min'), ('price','price_yuan_per_kwh')]:
            same(df[col], actual[source])
        up = np.maximum(df.q-df.g,0)
        down = np.maximum(df.g-df.q,0)
        contract = df.price*(df.g + 1.5*up + (.5 if settlement == 'N0' else -.5)*down)
        cost = contract + 5*df.price*df.e
        same(cost.to_numpy().reshape(334,144).sum(1), daily.cost)
        for col, vals in [('plan_cost',df.price*df.g),('up_cost',1.5*df.price*up),('down_fee',.5*df.price*down),('refund',df.price*down if settlement=='N1' else down*0),('emergency_cost',5*df.price*df.e)]:
            same(vals.to_numpy().reshape(334,144).sum(1), daily[col])
        summary = json.loads((folder / f'{mode}.json').read_text())
        same([cost.sum()], [summary['cost']])
        energy_residual = float((df.actual_grid+df.pv-df.load-df.spill-.1*df.c-(1/.9-1)*df.d).sum()-(df.E_end.iloc[-1]-df.E_start.iloc[0]))
        assert abs(energy_residual)<1e-6
        bookname = f'result4-{mode[-1]}.xlsx'
        wb = load_workbook(folder/bookname, data_only=True)
        errors=[]
        for sheet, values, fees in [('计划购电量',df.g,df.price*df.g)] + ([('调整购电量',df.q,contract)] if mode=='q4-3' else []):
            ws=wb[sheet]
            rows=list(ws.iter_rows(min_row=2,max_row=335,max_col=147,values_only=True))
            assert [pd.Timestamp(row[0]).strftime('%Y-%m-%d') for row in rows]==daily.date.tolist()
            mat=values.to_numpy().reshape(334,144)
            expected=np.column_stack([mat,mat.sum(1),fees.to_numpy().reshape(334,144).sum(1)])
            errors.append(same([r[1:] for r in rows],expected))
        battery=wb['充放电量']
        events=[]
        for di,(date,day) in enumerate(df.groupby('date',sort=False)):
            for block in range(6):
                row=2+di*6+block
                errors.append(same([battery.cell(row,3).value,battery.cell(row,4).value],day.iloc[block*24:(block+1)*24][['c','d']].sum()))
            assert pd.Timestamp(battery.cell(2+di*6,1).value).strftime('%Y-%m-%d')==date
            errors.append(same([battery.cell(2+di*6,6).value,battery.cell(3+di*6,6).value],[day.E_start.iloc[0],day.E_end.iloc[-1]]))
            e=day.e.to_numpy()
            active=e>1e-7
            indices=np.flatnonzero(active)
            chunks=np.split(indices,np.flatnonzero(np.diff(indices)>1)+1) if len(indices) else []
            def tm(t): return f'{t//6:02d}:{(t%6)*10:02d}'
            if not chunks: events.append((date,'无',0.))
            for chunk in chunks: events.append((date,f'{tm(int(chunk[0]))}-{tm(int(chunk[-1])+1)}',float(e[chunk].sum())))
        observed=[]; current=None
        for row in wb['紧急购电量'].iter_rows(min_row=2,max_col=3,values_only=True):
            if all(v is None for v in row): continue
            if row[0] is not None: current=pd.Timestamp(row[0]).strftime('%Y-%m-%d')
            observed.append((current,row[1],row[2]))
        assert len(observed)==len(events)
        assert [(x[0],x[1]) for x in observed]==[(x[0],x[1]) for x in events]
        errors.append(same([x[2] for x in observed],[x[2] for x in events]))
        root_identical=None
        if settlement=='N0':
            root_identical=(OUT/bookname).read_bytes()==(folder/bookname).read_bytes()
            assert root_identical
        wb.close()
        report[f'{settlement}/{mode}']={'cost':float(cost.sum()),'energy_residual_kwh':energy_residual,'max_workbook_numeric_error':max(errors),'emergency_rows':len(events),'main_copy_identical':root_identical}
        print(settlement,mode,report[f'{settlement}/{mode}'],flush=True)
(Path(__file__).parent/'artifact_checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
