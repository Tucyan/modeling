import sys,runpy
from pathlib import Path
if __name__=='__main__' and '--legacy' not in sys.argv and (Path(__file__).resolve().parents[1]/'outputs/q2_review_v2/selection_protocol.json').exists():
    runpy.run_path(str(Path(__file__).with_name('q2_review_verify.py')),run_name='__main__')
    raise SystemExit(0)
from q2_solver import ROOT,OUT,DATA,read_data,optimize
from q2_analyze import audit,intervals
import openpyxl, numpy as np,pandas as pd,json,hashlib,re

x=pd.read_csv(OUT/'dispatch_10min.csv'); audit(x)
p=json.loads((OUT/'workbook_payload.json').read_text(encoding='utf8'))
wb=openpyxl.load_workbook(OUT/'result2.xlsx',data_only=True)
assert wb.sheetnames==['计划购电量','充放电量','紧急购电量']
ws=wb.worksheets[0]
arr=np.array(list(ws.iter_rows(min_row=2,max_row=335,min_col=2,max_col=147,values_only=True)),float)
assert arr.shape==(334,146) and np.isfinite(arr).all()
assert np.max(abs(arr-np.array(p['plan'])))<1e-7
for j,date in enumerate(sorted(x.date.unique())):
    assert str(ws.cell(j+2,1).value)[:10]==date
assert ws.max_row==335 and ws.max_column==147
ws=wb.worksheets[1]
for j,row in enumerate(p['battery'],2):
    for col,value in enumerate(row,1):
        got=ws.cell(j,col).value
        if value is None: assert got is None,(j,col,got)
        elif col==1: assert str(got)[:10]==value
        elif isinstance(value,(float,int)): assert abs(got-value)<1e-7
        else: assert got==value,(j,col,got,value)
assert ws.max_row==2005
ws=wb.worksheets[2]; events=0; total=0.
for j,row in enumerate(p['emergency'],2):
    for col,value in enumerate(row,1):
        got=ws.cell(j,col).value
        if value is None: assert got is None
        elif col==1: assert str(got)[:10]==value
        elif col==3: assert abs(got-value)<1e-7;total+=got
        else: assert got==value
    events+=row[2]>1e-7
assert abs(total-x.e.sum())<1e-6
assert ws.max_row==len(p['emergency'])+1
assert all(cell.data_type!='e' for s in wb for row in s for cell in row)

# Reconstruct every delivered day from information available at that midnight.
_,L,V,price=read_data(); f=np.load(OUT/'forecasts.npz')['P0']; pn=f[:,0]-f[:,1]
actual=L-V; gaps=[]; max_purchase_error=0
for i in range(31,365):
    history=np.arange(max(7,i-56),i)
    net=pn[i]+np.quantile(actual[history]-pn[history],.8,axis=0)
    rows=x.iloc[(i-31)*144:(i-30)*144]
    opt=optimize(net,price,float(rows.E_start.iloc[0]))
    max_purchase_error=max(max_purchase_error,float(np.max(abs(opt['g']-rows.g))))
    gaps.append(opt['gap'])
assert max_purchase_error<1e-6
assert max(gaps)<1e-5
paper=ROOT/'outputs/q2_review_v2/第二问原版论文_存档.md'
if not paper.exists(): paper=ROOT/'docs/result/第二问求解与结果分析.md'
text=paper.read_text(encoding='utf8')
assert text.count('$$')%2==0
for date,day in p['specified'].items():
    assert date in text
    section=text.split('### '+date,1)[1].split('### ',1)[0]
    for key in ['plan_kwh','plan_cost','emergency_kwh','emergency_cost','total_cost','initial','terminal']:
        assert f'{day[key]:,.2f}' in section,(date,key)
for line in text.splitlines():
    if '](' not in line: continue
    for target in re.findall(r'\]\(([^)]+)\)', line):
        target = target.split('#', 1)[0]
        if not target or target.startswith(('http://', 'https://', 'mailto:')):
            continue
        path = (paper.parent / target).resolve()
        assert not Path(target).is_absolute(), target
        assert path.exists(), target
result={'plan_cells_checked':334*146,'battery_days_checked':334,'emergency_event_count':events,
  'emergency_total_kwh':total,'paper_specified_dates':4,'midnight_reproduced_days':334,
  'maximum_purchase_reproduction_error_kwh':max_purchase_error,'maximum_lp_primal_dual_gap_yuan':max(gaps),
  'xlsx_sha256':hashlib.sha256((OUT/'result2.xlsx').read_bytes()).hexdigest(),
  'solver_source_sha256':hashlib.sha256((ROOT/'src/q2_solver.py').read_bytes()).hexdigest(),
  'paper_sha256':hashlib.sha256(paper.read_bytes()).hexdigest()}
(OUT/'workbook_verification.json').write_text(json.dumps(result,indent=2),encoding='utf8')
print(json.dumps(result,indent=2))
